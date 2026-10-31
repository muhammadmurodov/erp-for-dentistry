"""Importing a bank statement from CSV or Excel.

The bank's export format is not known (docs/OPEN_QUESTIONS.md question B3),
so the column names are **configuration**: a `column_mapping` on the bank
account maps this system's field names onto whatever headings the file uses.
Supporting a different bank, or a changed export, is then a settings change
rather than a code change — which is also what "design an interface that
allows bank integrations to be added later" asks for.

Nothing imported here posts to the ledger. A statement line is evidence that
money moved, not the event that moved it; the accounting consequence belongs
to the match (see `reconciliation.py`).
"""

from __future__ import annotations

import csv
import hashlib
import io
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import NotFound, ValidationFailed
from app.core.money import ZERO, quantize, to_amount
from app.modules.audit import events
from app.modules.audit.service import AuditRecorder
from app.modules.banking.enums import ImportStatus, TransactionDirection
from app.modules.banking.models import (
    BankAccount,
    BankStatementImport,
    BankTransaction,
)

#: The fields this importer understands, and the headings it will look for
#: when an account has no mapping of its own. The Cyrillic names are the ones
#: the clinic's own workbook uses, so an export shaped like that sheet works
#: without configuration.
DEFAULT_COLUMN_MAPPING: dict[str, list[str]] = {
    "value_date": ["value_date", "date", "ДатаОперации", "Дата", "Сана"],
    "credit": ["credit", "amount_in", "Поступление", "Кирим", "Приход"],
    "debit": ["debit", "amount_out", "Расход", "Чиқим", "Расходы"],
    "description": [
        "description",
        "purpose",
        "Назначение платежа",
        "Тўлов мақсади",
    ],
    "counterparty_name": [
        "counterparty",
        "counterparty_name",
        "Наименование корреспондента",
    ],
    "counterparty_account": ["counterparty_account", "Счёт корреспондента"],
    "counterparty_mfo": ["counterparty_mfo", "МФО банка корреспондента", "МФО"],
    "document_number": ["document_number", "№ Документа", "Номер документа"],
}

#: Date formats seen in Uzbek bank exports, tried in order.
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y")

MAX_ROWS = 20_000

#: Footer labels seen at the foot of an export's table, in the three
#: languages the statements arrive in. A dateless row whose first cell says
#: one of these is the statement's own total, not a line the mapping failed
#: on, and reporting it as a problem on every single import is how a problem
#: list gets ignored. Anything else without a date is still reported —
#: guessing is how a third of a statement goes missing unnoticed.
TOTALS_LABELS = frozenset(
    {
        "итого",
        "всего",
        "оборот",
        "обороты",
        "итого за период",
        "jami",
        "jami:",
        "hammasi",
        "aylanma",
        "total",
        "totals",
    }
)

#: Separators that appear inside amount and header cells in bank exports.
#: Written as escapes because an invisible character in a string literal is
#: unreviewable — and the linter is right to object to one.
NO_BREAK_SPACE = "\u00a0"
NARROW_NO_BREAK_SPACE = "\u202f"
THIN_SPACE = "\u2009"
GROUPING_CHARACTERS = (
    NO_BREAK_SPACE,
    NARROW_NO_BREAK_SPACE,
    THIN_SPACE,
    "'",
    " ",
)


@dataclass(frozen=True, slots=True)
class ParsedRow:
    value_date: date
    direction: TransactionDirection
    amount: Decimal
    description: str | None
    counterparty_name: str | None
    counterparty_account: str | None
    counterparty_mfo: str | None
    document_number: str | None
    raw: dict[str, Any]


@dataclass(slots=True)
class ImportResult:
    import_id: uuid.UUID
    rows_read: int = 0
    rows_imported: int = 0
    duplicates_skipped: int = 0
    rows_rejected: int = 0
    problems: list[str] = field(default_factory=list)
    period_from: date | None = None
    period_to: date | None = None


def fingerprint_for(
    *,
    value_date: date,
    direction: TransactionDirection,
    amount: Decimal,
    document_number: str | None,
    description: str | None,
) -> str:
    """A digest of what a bank cannot change between exports.

    Re-downloading a statement is routine, and an importer that creates a
    second copy of every line each time is worse than useless. These five
    fields identify a line; the running balance and row number do not,
    because they shift when a date range changes.
    """
    parts = [
        value_date.isoformat(),
        direction.value,
        str(quantize(amount)),
        (document_number or "").strip(),
        " ".join((description or "").split()).lower(),
    ]
    return hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()


def _delimiter_for(text: str) -> str:
    """Pick the separator by how well it explains the file.

    `csv.Sniffer` is no use here: it inspects the opening characters, and a
    bank statement opens with a title block, the account holder and a blank
    line rather than a table, so it reports that it cannot tell — and the
    fallback to a comma then reads a semicolon-separated statement as one
    very wide column.

    Each candidate is tried instead and the one that produces the widest
    table that most rows agree on wins, which is what makes a file tabular
    in the first place.
    """
    lines = [line for line in text.splitlines()[:60] if line.strip()]
    if not lines:
        return ","
    best, best_score = ",", (0, 0)
    for candidate in (";", ",", "\t", "|"):
        widths = [len(row) for row in csv.reader(lines, delimiter=candidate)]
        modal = max(set(widths), key=widths.count)
        score = (modal, widths.count(modal))
        if modal > 1 and score > best_score:
            best, best_score = candidate, score
    return best


def _normalise_header(value: str) -> str:
    return " ".join(str(value).replace(NO_BREAK_SPACE, " ").split()).strip().lower()


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    # Excel often carries a time component the bank never meant.
    text = text.split(" ")[0].split("T")[0]
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_amount(value: Any) -> Decimal | None:
    """Read a money cell without ever going through a binary float.

    Exports arrive with thousands separators as spaces, non-breaking spaces
    or apostrophes, and with either a comma or a dot for the decimal mark.
    """
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        # openpyxl hands back floats. Going via `repr` keeps the shortest
        # representation that round-trips, which is the number the file
        # displayed rather than its binary expansion.
        return Decimal(repr(value))
    text = str(value)
    for character in GROUPING_CHARACTERS:
        text = text.replace(character, "")
    text = text.strip()

    if not text:
        return None
    # A comma is the decimal mark when there is no dot, and a thousands
    # separator when there is. Either convention appears in Uzbek exports.
    decimal_comma = "," in text and "." not in text
    text = text.replace(",", "." if decimal_comma else "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


class StatementImporter:
    def __init__(self, session: AsyncSession, audit: AuditRecorder) -> None:
        self._session = session
        self._audit = audit

    # ------------------------------------------------------------- mapping

    @staticmethod
    def resolve_columns(headers: list[str], mapping: dict[str, Any] | None) -> dict[str, int]:
        """Work out which column holds which field.

        An account's own mapping wins; otherwise the known aliases are tried.
        Returns field name → column index.
        """
        normalised = {_normalise_header(h): i for i, h in enumerate(headers) if h}
        resolved: dict[str, int] = {}

        for field_name, aliases in DEFAULT_COLUMN_MAPPING.items():
            configured = (mapping or {}).get(field_name)
            candidates = [configured] if configured else list(aliases)
            for candidate in candidates:
                index = normalised.get(_normalise_header(str(candidate)))
                if index is not None:
                    resolved[field_name] = index
                    break
        return resolved

    # -------------------------------------------------------------- import

    async def import_statement(
        self,
        *,
        organization_id: uuid.UUID,
        bank_account_id: uuid.UUID,
        filename: str,
        content: bytes,
        actor_user_id: uuid.UUID | None = None,
    ) -> ImportResult:
        account = await self._get_account(organization_id, bank_account_id)
        rows = self._read(filename, content)
        if not rows:
            raise ValidationFailed("The file contains no rows")
        if len(rows) > MAX_ROWS:
            raise ValidationFailed(
                f"The file has {len(rows)} rows; the limit is {MAX_ROWS}",
                details={"rows": len(rows)},
            )

        header_index = self._find_header_row(rows, account.column_mapping)
        if header_index is None:
            raise ValidationFailed(
                "Could not find a header row with a date column and an amount "
                "column. Configure the account's column mapping to match this "
                "bank's export.",
                details={
                    "expected_any_of": DEFAULT_COLUMN_MAPPING["value_date"],
                    "first_row": [str(c) for c in rows[0][:12]],
                },
            )

        headers = [str(c) if c is not None else "" for c in rows[header_index]]
        columns = self.resolve_columns(headers, account.column_mapping)
        missing = [f for f in ("value_date",) if f not in columns]
        if missing or not ({"credit", "debit"} & set(columns)):
            raise ValidationFailed(
                "The statement is missing a date column or any amount column",
                details={"found": sorted(columns), "headers": headers[:20]},
            )

        record = BankStatementImport(
            organization_id=organization_id,
            bank_account_id=bank_account_id,
            filename=filename[:255],
            status=ImportStatus.COMPLETED,
            imported_by_user_id=actor_user_id,
        )
        self._session.add(record)
        await self._session.flush()

        result = ImportResult(import_id=record.id)
        existing = await self._existing_fingerprints(organization_id, bank_account_id)
        seen_in_file: set[str] = set()

        for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            result.rows_read += 1
            parsed, problem = self._parse_row(row, columns, headers)
            if parsed is None:
                if problem:
                    result.rows_rejected += 1
                    result.problems.append(f"row {offset}: {problem}")
                continue

            digest = fingerprint_for(
                value_date=parsed.value_date,
                direction=parsed.direction,
                amount=parsed.amount,
                document_number=parsed.document_number,
                description=parsed.description,
            )
            if digest in existing or digest in seen_in_file:
                result.duplicates_skipped += 1
                continue
            seen_in_file.add(digest)

            self._session.add(
                BankTransaction(
                    organization_id=organization_id,
                    bank_account_id=bank_account_id,
                    import_id=record.id,
                    value_date=parsed.value_date,
                    direction=parsed.direction,
                    amount=parsed.amount,
                    currency=account.currency,
                    description=parsed.description,
                    counterparty_name=parsed.counterparty_name,
                    counterparty_account=parsed.counterparty_account,
                    counterparty_mfo=parsed.counterparty_mfo,
                    document_number=parsed.document_number,
                    fingerprint=digest,
                    raw=parsed.raw,
                )
            )
            result.rows_imported += 1
            if result.period_from is None or parsed.value_date < result.period_from:
                result.period_from = parsed.value_date
            if result.period_to is None or parsed.value_date > result.period_to:
                result.period_to = parsed.value_date

        record.rows_read = result.rows_read
        record.rows_imported = result.rows_imported
        record.duplicates_skipped = result.duplicates_skipped
        record.rows_rejected = result.rows_rejected
        record.problems = result.problems[:100] or None
        record.period_from = result.period_from
        record.period_to = result.period_to
        await self._session.flush()

        self._audit.record(
            events.STATEMENT_IMPORTED,
            organization_id=organization_id,
            target_type="bank_statement_import",
            target_id=record.id,
            target_label=filename[:255],
            business_event=f"bank_import:{record.id}",
            new_values={
                "bank_account": account.code,
                "rows_read": result.rows_read,
                "rows_imported": result.rows_imported,
                "duplicates_skipped": result.duplicates_skipped,
                "rows_rejected": result.rows_rejected,
                "period_from": result.period_from,
                "period_to": result.period_to,
            },
        )
        return result

    # ------------------------------------------------------------ internals

    async def list_imports(
        self,
        *,
        organization_id: uuid.UUID,
        bank_account_id: uuid.UUID | None = None,
        limit: int = 50,
    ) -> list[BankStatementImport]:
        """Every import run, including the ones that rejected rows.

        Kept visible on purpose: an import that silently dropped a third of
        its lines is the kind of thing that only surfaces at year end.
        """
        stmt = select(BankStatementImport).where(
            BankStatementImport.organization_id == organization_id
        )
        if bank_account_id is not None:
            stmt = stmt.where(BankStatementImport.bank_account_id == bank_account_id)
        stmt = stmt.order_by(BankStatementImport.created_at.desc()).limit(limit)
        return list((await self._session.execute(stmt)).scalars().all())

    async def _get_account(
        self, organization_id: uuid.UUID, bank_account_id: uuid.UUID
    ) -> BankAccount:
        stmt = select(BankAccount).where(
            BankAccount.organization_id == organization_id,
            BankAccount.id == bank_account_id,
        )
        account = (await self._session.execute(stmt)).scalar_one_or_none()
        if account is None:
            raise NotFound("Bank account not found")
        return account

    async def _existing_fingerprints(
        self, organization_id: uuid.UUID, bank_account_id: uuid.UUID
    ) -> set[str]:
        stmt = select(BankTransaction.fingerprint).where(
            BankTransaction.organization_id == organization_id,
            BankTransaction.bank_account_id == bank_account_id,
        )
        return set((await self._session.execute(stmt)).scalars().all())

    @staticmethod
    def _read(filename: str, content: bytes) -> list[list[Any]]:
        lowered = filename.lower()
        if lowered.endswith((".xlsx", ".xlsm")):
            import openpyxl

            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            rows = [list(r) for r in sheet.iter_rows(max_col=40, values_only=True)]
            workbook.close()
            return rows
        if lowered.endswith(".csv"):
            # Banks export in whatever encoding they like; try the common
            # ones rather than failing on a byte.
            for encoding in ("utf-8-sig", "utf-8", "cp1251"):
                try:
                    text = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValidationFailed(
                    "The file's text encoding could not be read. Save it as UTF-8 and try again."
                )
            delimiter = _delimiter_for(text)
            return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
        raise ValidationFailed(
            "Only .csv, .xlsx and .xlsm statements can be imported",
            details={"filename": filename},
        )

    @classmethod
    def _find_header_row(cls, rows: list[list[Any]], mapping: dict[str, Any] | None) -> int | None:
        """Find the header, which is rarely the first row.

        Bank exports routinely carry a title block, the account holder's
        details and a blank line before the table starts.
        """
        for index, row in enumerate(rows[:30]):
            headers = [str(c) if c is not None else "" for c in row]
            columns = cls.resolve_columns(headers, mapping)
            if "value_date" in columns and ({"credit", "debit"} & set(columns)):
                return index
        return None

    @staticmethod
    def _parse_row(
        row: list[Any], columns: dict[str, int], headers: list[str]
    ) -> tuple[ParsedRow | None, str | None]:
        def cell(name: str) -> Any:
            index = columns.get(name)
            if index is None or index >= len(row):
                return None
            return row[index]

        value_date = _parse_date(cell("value_date"))
        credit = _parse_amount(cell("credit"))
        debit = _parse_amount(cell("debit"))

        if value_date is None and credit is None and debit is None:
            # A blank separator row, or the totals line. Not an error.
            return None, None
        if value_date is None:
            first = next((c for c in row if c not in (None, "")), "")
            if _normalise_header(str(first)).strip(" :.-") in TOTALS_LABELS:
                return None, None  # the statement's own totals line
            return None, "no readable date"

        credit = credit or ZERO
        debit = debit or ZERO
        if credit > ZERO and debit > ZERO:
            return None, "both an inflow and an outflow on one line"
        if credit == ZERO and debit == ZERO:
            return None, None  # a zero line carries no information

        direction = TransactionDirection.IN if credit > ZERO else TransactionDirection.OUT
        amount = quantize(to_amount(abs(credit if credit > ZERO else debit)))
        if amount <= ZERO:
            return None, "amount is not positive"

        def text(name: str) -> str | None:
            value = cell(name)
            if value in (None, ""):
                return None
            return " ".join(str(value).split())[:255] or None

        raw = {
            headers[i][:64]: (str(row[i])[:256] if row[i] is not None else None)
            for i in range(min(len(headers), len(row)))
            if headers[i]
        }
        return (
            ParsedRow(
                value_date=value_date,
                direction=direction,
                amount=amount,
                description=text("description"),
                counterparty_name=text("counterparty_name"),
                counterparty_account=text("counterparty_account"),
                counterparty_mfo=text("counterparty_mfo"),
                document_number=text("document_number"),
                raw=raw,
            ),
            None,
        )
